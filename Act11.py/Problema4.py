from openpyxl import load_workbook
Workbook = load_workbook("mi_libro.xlsx")
hoja = Workbook["Ventas"]
celda = hoja.cell(3, 3)
celda.value = 0.55
Workbook.save("mi_libro.xlsx")