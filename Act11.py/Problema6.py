from openpyxl import load_workbook
wk = load_workbook("mi_libro.xlsx")
Inventario = wk.create_sheet("Inventario")

inventario =  wk["Inventario"]
lxl = [
["Producto", "Stock"],
["Manzanas", 100],
["Naranjas", 80]
 ]
for i, fila in enumerate(lxl):
    for j, valor in enumerate(fila):
        inventario.cell(row = i + 1, column = j + 1, value = valor)
wk.save("mi_libro.xlsx")
Ventas = wk["Ventas"]

lista_ventas = []
for i in Ventas.iter_rows():
    lista = [celda.value for celda in i]
    lista_ventas.append(lista)
dic ={ "Ventas": lista_ventas,
      "Inventario": lxl}
print(dic)