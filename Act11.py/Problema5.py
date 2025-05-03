from openpyxl import load_workbook
wk = load_workbook("mi_libro.xlsx")
hoja = wk["Ventas"]
hoja.cell(1, 4, "Total")
for i in range(2, 4):
    hoja.cell(i, 4).value = f"=B{i}*c{i}"
wk.save("mi_libro.xlsx")