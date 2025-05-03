from openpyxl import Workbook
libro = Workbook()
pag1 = libro.create_sheet("pag1")
pag1.title = "Datos"
print(libro.sheetnames)
libro.save("mi_libro.xlsx")