from openpyxl import Workbook
libro = Workbook()
pag = libro.create_sheet("pag")
pag.title = "Ventas"
Ventas = libro["Ventas"]
Ventas.cell(1, 1, "Producto")
Ventas.cell(1, 2, "Cantidad")
Ventas.cell(1, 3, "Precio")
Ventas.cell(2, 1, "Naranjas")
Ventas.cell(2, 2, 50)
Ventas.cell(2, 3, 0.5)
Ventas.cell(3, 1, "Manzanas")
Ventas.cell(3, 2, 30)
Ventas.cell(3, 3, 0.75)
libro.save("mi_libro.xlsx")