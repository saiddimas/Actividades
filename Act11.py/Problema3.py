from openpyxl import load_workbook
workbook = load_workbook("mi_libro.xlsx")
hoja = workbook["Ventas"]
lista = []

for i in hoja.iter_rows():
    filas_listas = [celda.value for celda in i]
    lista.append(filas_listas)
print(lista)
